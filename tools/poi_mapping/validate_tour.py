# -*- coding: utf-8 -*-
"""
Valida os hotspots de uma visita contra o indice OpenSearch, corrige o que e
mecanico (com --fix) e gera, em FIXES/, o Excel das pecas que ficam por
confirmar (imagem, nome, ficha, sala, panorama, numero e uma nota simples para
quem escreve os numeros no 3DVista).

Aceita o projecto 3DVista (.vtp) ou a pasta exportada (com script_general.js):

    python tools/poi_mapping/validate_tour.py "Visitas/MMC/MMC.vtp" --museum mnc --index cultural_heritage_artifacts_v4
    python tools/poi_mapping/validate_tour.py tmp/visitas/mnaz --museum mnaz --fix
    python tools/poi_mapping/validate_tour.py tmp/visitas/mnaz --museum mnaz --images-dir renders/ --out-dir FIXES

--fix corrige directamente no ficheiro (com copia .bak ao lado):
  - numeros que existem mas estao escritos de outra forma -> como no indice
  - varios numeros colados numa string -> um elemento por peca
  - vitrines «SN» -> <MUSEU>_SN_<n>, com o numero que esta no titulo
  - aspas curvas -> direitas; HTML no titulo -> texto; location com espacos -> underscores
  - URLs de janela com maiusculas trocadas ou quebras de linha
O que o script nao consegue resolver (numeros que nao existem no Raiz) fica no Excel.

Precisa do backend (le o .env para chegar ao OpenSearch). Sai com codigo 1 se
houver numeros que nao existem.
"""
from __future__ import annotations

import argparse
import datetime as dt
import glob
import html as htmlmod
import json
import os
import re
import shutil
import sys
import unicodedata
import warnings
import zipfile
from pathlib import Path

HERE = Path(__file__).resolve().parent
REPO = HERE.parents[1]
BACKEND = REPO  # raiz do repositório do backend
sys.path.insert(0, str(HERE))
sys.path.insert(0, str(BACKEND))
warnings.filterwarnings("ignore")

SN_OK_RE = re.compile(r"^[A-Za-z]+_SN_\d+$")
SN_RAW_RE = re.compile(r"(^|[_ ])SN$", re.I)
SN_NAMED_RE = re.compile(r"^([A-Za-z]+)[ _]SN[ _]([A-Za-z]\w*)$")   # «MNAz_SN_Senhora»: SN com nome em vez de numero
IMG_EXT = (".jpg", ".jpeg", ".png", ".webp")
STR = r'"((?:[^"\\]|\\.)*)"'


def connect():
    """Cliente OpenSearch + indice por defeito.

    1) Se o backend do assistente estiver ao lado (repo completo), usa-o e le o
       .env da raiz do repositório, como a app.
    2) Senao (pasta de ferramentas sozinha, noutra maquina), le um .env na pasta
       deste script ou as variaveis de ambiente: OPENSEARCH_HOST, OPENSEARCH_PORT,
       OPENSEARCH_USERNAME, OPENSEARCH_PASSWORD, OPENSEARCH_VERIFY_CERTS,
       OPENSEARCH_INDEX_ARTIFACT. Precisa do pacote opensearch-py.
    """
    if (BACKEND / "app" / "core" / "config.py").exists():
        os.chdir(BACKEND)
        from app.core.config import get_settings
        from app.services.opensearch_client import get_opensearch_gateway
        return get_opensearch_gateway()._ensure_client(), get_settings().OPENSEARCH_INDEX_ARTIFACT

    env = dict(os.environ)
    env_file = HERE / ".env"
    if env_file.exists():
        for line in env_file.read_text(encoding="utf-8").splitlines():
            line = line.strip()
            if line and not line.startswith("#") and "=" in line:
                k, v = line.split("=", 1)
                env.setdefault(k.strip(), v.strip().strip('"').strip("'"))
    host = env.get("OPENSEARCH_HOST", "").strip()
    if not host:
        sys.exit("Sem backend ao lado e sem OPENSEARCH_HOST: preenche tools/poi_mapping/.env (ver .env.example).")
    from urllib.parse import urlparse
    from opensearchpy import OpenSearch
    parsed = urlparse(host if "://" in host else f"{env.get('OPENSEARCH_SCHEME', 'https')}://{host}")
    port = int(env.get("OPENSEARCH_PORT") or parsed.port or (443 if parsed.scheme == "https" else 80))
    user, pwd = env.get("OPENSEARCH_USERNAME"), env.get("OPENSEARCH_PASSWORD")
    verify = env.get("OPENSEARCH_VERIFY_CERTS", "false").lower() in ("1", "true", "yes")
    client = OpenSearch(
        hosts=[{"host": parsed.hostname, "port": port, "scheme": parsed.scheme}],
        http_auth=(user, pwd) if user and pwd else None,
        use_ssl=parsed.scheme == "https", verify_certs=verify, ssl_show_warn=False, timeout=60,
    )
    return client, env.get("OPENSEARCH_INDEX_ARTIFACT", "cultural_heritage_artifacts")


def norm(s):
    return re.sub(r"[^A-Z0-9]", "", (s or "").upper())


def norm_title(s):
    s = unicodedata.normalize("NFKD", s or "").encode("ascii", "ignore").decode()
    return re.sub(r"[^a-z0-9]+", " ", s.lower()).strip()


def clean_html(s):
    s = re.sub(r"<br\s*/?>", " · ", s or "")
    return re.sub(r"\s+", " ", htmlmod.unescape(re.sub(r"<[^>]+>", "", s))).strip(" ·")


def js(s):
    return json.dumps(s, ensure_ascii=False)


def new_code(title, ids, loc):
    arr = ", ".join(f'"{i}"' for i in ids)
    return f'window.postMessage({{ \n  type: "openedInfoWindow", \n  inventoryIds: [{arr}], \n  title: "{title}", \n  location: "{loc}" \n}}, "*") \n'


def sentence(code):
    return "try{eval('" + code.replace("\n", "").replace("\\", "\\\\").replace("'", "\\'") + "')}catch(e){console.log(e)}"


def backup(path: Path):
    b = path.with_name(path.name + ".bak-" + dt.datetime.now().strftime("%Y%m%d-%H%M"))
    shutil.copy2(path, b)
    return b


# ------------------------------------------------------------ leitura da visita
def hotspots_from_vtp(vtp: Path):
    from audit_vtp import load_project, audit
    conteudos = next((vtp.parent / n for n in ("Conteudos", "Conteúdos") if (vtp.parent / n).exists()), vtp.parent / "Conteudos")
    hotspots, pages, locales = audit(load_project(vtp), conteudos if conteudos.exists() else None)
    img_dir = conteudos / "HTML" / "iFrames" / "media" / "images"

    def ficha(page):
        if not page:
            return {}
        p = conteudos / "HTML" / (page.strip().replace("generated-html", "data")[:-5] + ".json")
        return json.load(open(p, encoding="utf-8")) if p.exists() else {}

    def media_names(f):
        return [(m.get("filename") if isinstance(m, dict) else m) or "" for m in f.get("media", [])]

    def image(f):
        return next((str(img_dir / n) for n in media_names(f) if n.lower().endswith(IMG_EXT) and (img_dir / n).exists()), None)

    out = []
    for h in hotspots:
        if not h["js"]:
            continue
        f = ficha(h["page"])
        pk = h["panoramaKey"] or ""
        code, sents = h["js"].get("code"), h["js"].get("sentences")
        stale = bool(code) and (sents is None or sents != sentence(code))
        out.append({"key": h["areaId"], "panorama": h["panorama"], "panoramaKey": pk,
                    "label": locales.get(pk + ".label"), "subtitle": locales.get(pk + ".subtitle"),
                    "title": h["js"]["title"],
                    "inventoryIds": h["js"]["inventoryIds"], "location": h["js"]["location"],
                    "smartQuotes": h["js"]["smartQuotes"], "hasHotspot": True, "sentencesStale": stale,
                    "ficha": clean_html(f.get("denomination", "")), "media": media_names(f), "image": image(f)})
    used = {h["page"].strip() for h in hotspots if h["page"]}
    for p, d in pages.items():
        if p in used:
            continue
        f = ficha(p)
        # «MNAA, inv. 297 Mov» -> um numero; «MNA 999.143.1, MNA 999.143.2» -> dois
        ids = [x.strip() for x in re.split(r"[,;]", d.get("inventory_number") or "") if re.search(r"\d", x)]
        out.append({"key": None, "panorama": None, "title": d.get("header"), "inventoryIds": ids,
                    "location": d.get("room"), "smartQuotes": False, "hasHotspot": False,
                    "ficha": clean_html(f.get("denomination", "")), "media": media_names(f), "image": image(f)})
    return out


def export_index(folder: Path):
    """Texto do script exportado, spans de objectos e locales."""
    from extract_poi_mapping import object_spans, innermost
    src = next((folder / n for n in ("script_general.js", "script_mobile.js") if (folder / n).exists()), None)
    if src is None:
        sys.exit(f"Nao encontrei script_general.js em {folder}")
    text = src.read_text(encoding="utf-8", errors="replace")
    spans = object_spans(text)
    starts = [s[0] for s in spans]
    locale = {}
    for lf in glob.glob(str(folder / "locale" / "*.txt")):
        for line in open(lf, encoding="utf-8", errors="replace"):
            if " = " in line:
                k, v = line.rstrip("\n").split(" = ", 1)
                locale.setdefault(k, v)
    return text, spans, starts, locale, innermost


def hotspots_from_export(folder: Path):
    from extract_poi_mapping import extract
    results, _, _ = extract(folder)
    text, spans, starts, locale, innermost = export_index(folder)

    def page_of(overlay_id):
        m = re.search(r'"id":"' + re.escape(overlay_id) + '"', text)
        if not m:
            return None
        s = innermost(spans, starts, m.start())
        areas = re.findall(r'"this\.((?:Flat)?HotspotPanoramaOverlayArea_[0-9A-F_]+)"', text[s[0]:s[1]])
        for area in areas + [overlay_id]:
            am = re.search(r'"id":"' + re.escape(area) + '"', text)
            if not am:
                continue
            a = innermost(spans, starts, am.start())
            k = re.findall(r"PopupWebFrameBehaviour_[0-9A-F_]+", text[a[0]:a[1]])
            if k and locale.get(k[0] + ".url"):
                return locale[k[0] + ".url"].strip()
        return None

    def local(rel):
        p = folder / rel
        if p.exists():
            return p
        low = rel.lower()
        for cand in glob.glob(str(folder / "**" / "*.html"), recursive=True):
            if os.path.relpath(cand, folder).replace(os.sep, "/").lower() == low:
                return Path(cand)
        return None

    # aspas curvas: ver o texto bruto da accao de cada hotspot
    smart, smart_obj = {}, {}
    for m in re.finditer("openedInfoWindow", text):
        s = innermost(spans, starts, m.start())
        obj = text[s[0]:s[1]]
        aid = re.search(r'"id":"([^"]+)"', obj)
        if not aid:
            continue
        key = aid.group(1)
        if not key.startswith(("overlay_", "sprite_")):
            pm = re.search(r'"this\.' + re.escape(key) + '"', text)
            if pm:
                ps = innermost(spans, starts, pm.start())
                pid = re.search(r'"id":"([^"]+)"', text[ps[0]:ps[1]])
                key = pid.group(1) if pid else key
        smart[key] = bool(re.search("[“”‘’]", obj))
        smart_obj[key] = obj

    out = []
    for r in results:
        r["smartQuotes"] = smart.get(r["overlayId"], False)
        if r["smartQuotes"] and r["title"] and re.search("[“”‘’]", r["title"]):
            # aspas curvas so dentro do titulo sao texto legitimo
            ttl = r["title"]
            r["smartQuotes"] = any(re.search("[“”‘’]", smart_obj[k].replace(ttl, ""))
                                   for k in [r["overlayId"]] if k in smart_obj)
        rel = page_of(r["overlayId"])
        ficha, image = "", None
        p = local(rel) if rel else None
        if p:
            page = p.read_text(encoding="utf-8", errors="replace")
            m = re.search(r'id="title-text">(.*?)</strong>', page, re.S)
            ficha = clean_html(m.group(1)) if m else ""
            imgs = re.findall(r"[^\"' (]+media/images/[^\"' )]+", page)
            if imgs:
                cand = (p.parent / imgs[0]).resolve()
                image = str(cand) if cand.exists() else None
        pk = r["panoramaKey"]
        out.append({"key": r["overlayId"], "panorama": locale.get(pk + ".label") or pk, "panoramaKey": pk,
                    "label": locale.get(pk + ".label"), "subtitle": locale.get(pk + ".subtitle"),
                    "title": r["title"],
                    "inventoryIds": r["inventoryIds"], "location": r["location"],
                    "smartQuotes": r.get("smartQuotes", False), "hasHotspot": True,
                    "ficha": ficha, "media": [], "image": image})
    return out


# ------------------------------------------------------------------- catalogo
def load_catalogue(client, index, museum):
    out, name = {}, None
    r = client.search(index=index, body={"size": 1000, "query": {"term": {"museum_id": museum}},
                                         "_source": ["inventory_number", "title", "museum"]}, scroll="2m")
    sid = r["_scroll_id"]
    while r["hits"]["hits"]:
        for h in r["hits"]["hits"]:
            src = h["_source"]
            name = name or src.get("museum")
            if src.get("inventory_number"):
                out[src["inventory_number"]] = src.get("title") or ""
        r = client.scroll(scroll_id=sid, scroll="2m")
    client.clear_scroll(scroll_id=sid)
    return out, name


PREFIX_RE = re.compile(r"^([A-Za-z]{2,5})[ ,_]*(Inv\.?)?[ ._:]*", re.I)


INV_MARK_RE = re.compile(r"^(?:[A-Za-z]{2,5}[ ,_]*)?Inv\.?[ ._:]*(?=.*\d)", re.I)


def strip_prefix(inv, dominant):
    """Tira um prefixo de museu («MMC_Inv.», «MNAA, inv. », «inv. ») quando tem o marcador
    Inv, ou quando e o prefixo que, nesta visita, se provou estar a mais."""
    s = inv.strip()
    m = INV_MARK_RE.match(s)
    if m and m.end() < len(s):
        return s[m.end():]
    if dominant:
        m = re.match(r"^" + re.escape(dominant) + r"[ ,_.:]*(?=.*\d)", s, re.I)
        if m and m.end() < len(s):
            return s[m.end():]
    return s


def lookup(inv, by_norm, dominant=None):
    """Numero como esta no indice, tolerando pontuacao e o prefixo de museu."""
    real = by_norm.get(norm(inv))
    if real:
        return real
    stripped = strip_prefix(inv, dominant)
    return by_norm.get(norm(stripped)) if stripped != inv.strip() else None


def find_dominant_prefix(ids, by_norm):
    """Prefixo (ex.: MMC) que so atrapalha: os numeros so resolvem depois de o tirar."""
    from collections import Counter
    c = Counter()
    for inv in ids:
        m = PREFIX_RE.match(inv.strip())
        if m and by_norm.get(norm(inv)) is None and by_norm.get(norm(inv.strip()[m.end():])):
            c[m.group(1).upper()] += 1
    top = c.most_common(1)
    return top[0][0] if top and top[0][1] >= 3 else None


def split_glued(inv, by_norm, dominant=None):
    """«MNAz_3097_Az_1000_Az» -> ['MNAz 3097 Az', 'MNAz 1000 Az']; «A3776-3784» expande o intervalo.
    Numeros com pontos («999.144.1») contam como um so. Devolve a lista se pelo menos
    duas partes existirem no indice; as que faltam ficam escritas no mesmo formato."""
    s = inv.strip()
    core = s if by_norm.get(norm(s)) else strip_prefix(s, dominant)
    groups = []
    for m in re.finditer(r"(\d+(?:\.\d+)*)(?:\s*-\s*(\d+))?", core):
        a, b = m.group(1), m.group(2)
        if b and a.isdigit() and int(b) > int(a) and int(b) - int(a) < 40:
            groups += [str(n) for n in range(int(a), int(b) + 1)]
        else:
            groups.append(a)
    if len(groups) < 2:
        return None
    prefix = re.match(r"^[A-Za-z]*", core).group(0)
    suffix = re.search(r"[A-Za-z][A-Za-z ]*$", core)
    suffix = suffix.group(0).strip() if suffix else ""
    cands = [f"{prefix} {g} {suffix}".strip() for g in groups]
    parts = [by_norm.get(norm(c)) or c for c in cands]
    found = sum(1 for c in cands if by_norm.get(norm(c)))
    return parts if found >= 2 else None


# ------------------------------------------------------------------ correccao
def plan_fixes(hotspots, catalogue, by_norm, museum, location_from=None):
    """Para cada hotspot, o que ficaria depois da correccao mecanica.
    location_from: 'subtitle' ou 'label' -> o location passa a ser esse texto do panorama."""
    underscores_ok = sum("_" in k for k in catalogue) > 0.05 * len(catalogue)
    # estilo do indice: «A 4121» (letras, espaco, digitos) ou «A4121»
    space_style = sum(bool(re.match(r"^[A-Za-z]+ \d", k)) for k in catalogue) > sum(bool(re.match(r"^[A-Za-z]+\.?\d", k)) for k in catalogue)
    locs = [h["location"] for h in hotspots if h["location"]]
    prefer_us = sum("_" in l for l in locs) >= sum(" " in l for l in locs)
    dominant = find_dominant_prefix([i for h in hotspots for i in h["inventoryIds"]], by_norm)
    plan = {}
    for h in hotspots:
        if not h["hasHotspot"]:
            continue
        ids, changes = [], []
        for raw in h["inventoryIds"]:
            tokens = [raw]
            # «A3986/MNA.999.144.1» sao duas pecas; «MNAz 227/1 Az» e um numero so.
            # So se parte quando o todo nao resolve e pelo menos uma parte resolve sozinha.
            whole_ok = raw in catalogue or lookup(raw, by_norm, dominant) or split_glued(raw, by_norm, dominant)
            if not whole_ok:
                parts = [t.strip() for t in re.split(r"\s*[/;]\s*", raw) if t.strip()]
                if len(parts) > 1 and any(lookup(t, by_norm, dominant) for t in parts):
                    tokens = parts
                    changes.append(f"{raw} -> {' | '.join(tokens)} (varias pecas)")
            for inv in tokens:
                if inv in catalogue or SN_OK_RE.match(inv):
                    ids.append(inv); continue
                nm = SN_NAMED_RE.match(inv)
                if nm:
                    # SN com nome («MNAz_SN_Senhora»): nao e um numero do Raiz nem segue a norma;
                    # so se normaliza a escrita — o numero de vitrine tem de ser atribuido a mao
                    fixed = f"{nm.group(1)}_SN_{nm.group(2)}"
                    ids.append(fixed)
                    if fixed != inv:
                        changes.append(f"{inv} -> {fixed}")
                    continue
                if SN_RAW_RE.search(inv):
                    m = re.match(r"(\d+)\.\s", h["title"] or "")
                    if m:
                        ids.append(f"{museum.upper()}_SN_{m.group(1)}"); changes.append(f"{inv} -> {ids[-1]}")
                    else:
                        ids.append(inv)
                    continue
                real = lookup(inv, by_norm, dominant)
                if real:
                    ids.append(real); changes.append(f"{inv} -> {real}"); continue
                parts = split_glued(inv, by_norm, dominant)
                if parts:
                    ids += parts; changes.append(f"{inv} -> {' | '.join(parts)}"); continue
                fixed = strip_prefix(inv, dominant)
                fixed = re.sub(r"\s+", " ", fixed.replace("_", " ")).strip() if not underscores_ok else fixed
                if space_style:
                    fixed = re.sub(r"^([A-Za-z]+)[ .]*(\d)", r"\1 \2", fixed)
                ids.append(fixed)
                if fixed != inv:
                    changes.append(f"{inv} -> {fixed} (formato; continua sem existir)")
        title = clean_html(h["title"]).replace('"', "'")
        if title != (h["title"] or ""):
            changes.append("titulo sem HTML")
        loc = (h["location"] or "").strip()
        if location_from:
            wanted = (h.get(location_from) or h.get("subtitle") or h.get("label") or "").strip()
            if wanted and wanted != loc:
                loc = wanted; changes.append(f"location -> {loc}")
        elif prefer_us and " " in loc:
            loc = loc.replace(" ", "_"); changes.append(f"location -> {loc}")
        if h["smartQuotes"] and not re.search(r"[“”]", h["title"] or ""):
            changes.append("aspas curvas -> direitas")
        if h.get("sentencesStale"):
            changes.append("sentences desatualizado (o 3DVista exportava a versao antiga)")
        if changes:
            plan[h["key"]] = {"ids": ids, "title": title, "location": loc, "changes": changes}
    return plan


def apply_fix_vtp(vtp: Path, plan):
    from audit_vtp import object_spans, innermost
    with zipfile.ZipFile(vtp) as z:
        text = z.read("script.js").decode("utf-8")
        entries = [(i, z.read(i.filename)) for i in z.infolist()]
    spans, starts = object_spans(text)
    edits = []
    for area_id, p in plan.items():
        m = re.search(r'"id":"' + re.escape(area_id) + '"', text)
        if not m:
            continue
        span = innermost(spans, starts, m.start())   # o objecto da area, seja qual for a ordem das chaves
        if not span:
            continue
        seg = text[span[0]:span[1]]
        cm = re.search(r'"code":' + STR, seg)
        if not cm:
            continue
        code = new_code(p["title"], p["ids"], p["location"])
        new_sent = '"sentences":[' + js(sentence(code)) + "]"
        # o «sentences» certo e o que contem o openedInfoWindow (a area pode ter outros behaviours)
        sm = next((x for x in re.finditer(r'"sentences":\[' + STR + r"\]", seg) if "openedInfoWindow" in x.group(1)), None)
        seg2 = seg.replace(cm.group(0), '"code":' + js(code))
        if sm:
            seg2 = seg2.replace(sm.group(0), new_sent)
        else:
            seg2 = seg2.replace('"code":' + js(code), new_sent + ',"code":' + js(code))
        if seg2 != seg:
            edits.append((span[0], span[1], seg2))
    n = len(edits)
    for a, b, seg2 in sorted(edits, reverse=True):
        text = text[:a] + seg2 + text[b:]
    urls = 0
    from urllib.parse import unquote
    html_root = next((vtp.parent / n / "HTML" for n in ("Conteudos", "Conteúdos") if (vtp.parent / n / "HTML").exists()), None)
    for m in list(re.finditer(r'"(PopupWebFrameBehaviour_[0-9A-F_]+\.url)":\{"value":' + STR, text)):
        dec = json.loads('"' + m.group(2) + '"')
        fixed = dec.strip().replace("iframes/", "iFrames/")
        # «MNAA_297_%20Mesa.html»: o ficheiro foi renomeado sem o espaco; apontar para o que existe
        if html_root and "%20" in fixed and not (html_root / unquote(fixed)).exists():
            for cand in (fixed.replace("%20", ""), fixed.replace("%20", "_")):
                if (html_root / cand).exists():
                    fixed = cand
                    break
        if fixed != dec:
            text = text.replace(m.group(0), f'"{m.group(1)}":{{"value":' + js(fixed)); urls += 1
    if not n and not urls:
        return 0, 0, None
    tmp = vtp.with_suffix(".tmp")
    with zipfile.ZipFile(tmp, "w") as zout:
        for info, data in entries:
            zout.writestr(info, text.encode("utf-8") if info.filename == "script.js" else data, compress_type=info.compress_type)
    b = backup(vtp)
    try:
        os.replace(tmp, vtp)
    except PermissionError:
        tmp.unlink(missing_ok=True)
        b.unlink(missing_ok=True)
        sys.exit(f"Nao consegui gravar {vtp.name}: o projecto esta aberto no 3DVista. Fecha-o e volta a correr.")
    return n, urls, b


def apply_fix_export(folder: Path, plan):
    from extract_poi_mapping import object_spans, innermost
    total, backups = 0, []
    for name in ("script_general.js", "script_mobile.js"):
        f = folder / name
        if not f.exists():
            continue
        text = f.read_text(encoding="utf-8", errors="replace")
        spans = object_spans(text)
        starts = [s[0] for s in spans]

        def own_id(o):
            m = re.search(r'"id":"([^"]+)"', o)
            return m.group(1) if m else None

        def parent_overlay(child):
            for m in re.finditer(r'"this\.' + re.escape(child) + '"', text):
                s = innermost(spans, starts, m.start())
                pid = own_id(text[s[0]:s[1]])
                if pid and pid.startswith(("overlay_", "sprite_")):
                    return pid
            return None

        edits = []
        for m in re.finditer("openedInfoWindow", text):
            s = innermost(spans, starts, m.start())
            obj = text[s[0]:s[1]]
            aid = own_id(obj)
            if not aid:
                continue
            key = aid if aid.startswith(("overlay_", "sprite_")) else parent_overlay(aid)
            p = plan.get(key)
            cm = re.search(r'"click":' + STR, obj)
            if not p or not cm:
                continue
            click = json.loads('"' + cm.group(1) + '"')
            new_click = re.sub(r"try\{eval\('.*?'\)\}catch\(e\)\{console\.log\(e\)\}",
                               lambda _: sentence(new_code(p["title"], p["ids"], p["location"])), click, count=1, flags=re.S)
            if new_click != click:
                edits.append((s[0] + cm.start(), s[0] + cm.end(), '"click":' + js(new_click)))
        if edits:
            backups.append(backup(f))
            for a, b, rep in sorted(edits, reverse=True):
                text = text[:a] + rep + text[b:]
            f.write_text(text, encoding="utf-8")
            total += len(edits)
    return total, backups


# ---------------------------------------------------------------------- notas
def q(xs):
    return ", ".join("«" + x + "»" for x in xs)


def nota_para(h, missing, museum_name, by_title, by_norm):
    partes = []
    if not h["hasHotspot"]:
        partes.append("Esta ficha existe na visita mas ainda não tem hotspot em nenhum panorama.")
    onde = "no hotspot" if h["hasHotspot"] else "na ficha, antes de criar o hotspot"
    glued = {m: split_glued(m, by_norm) for m in missing}
    colados = [m for m in missing if glued[m]]
    sn_ids = [m for m in missing if m not in colados and (SN_OK_RE.match(m) or SN_NAMED_RE.match(m))]
    dep = [m for m in missing if m not in colados and m not in sn_ids and re.match(r"^(MNA\b|DEP)", m, re.I)]
    outros = [m for m in missing if m not in colados and m not in dep and m not in sn_ids]
    if sn_ids:
        pl = len(sn_ids) > 1
        partes.append(f"{'Os identificadores' if pl else 'O identificador'} {q(sn_ids)} {'são' if pl else 'é'} de peça sem número de inventário (SN). "
                      "Ainda não está no índice do assistente: tem de ser indexada à mão com este identificador. Não é preciso mexer no hotspot.")
    for m in colados:
        partes.append(f"«{m}» junta várias peças num só número. Separar em elementos, um por peça: {q(glued[m])}.")
    if dep:
        pl = len(dep) > 1
        partes.append(f"{'Os números' if pl else 'O número'} {q(dep)} {'são' if pl else 'é'} de uma peça em depósito, "
                      f"de outra instituição. O Raiz do {museum_name} não tem esta peça, por isso o assistente não a reconhece. "
                      "É preciso decidir com o museu se ela entra no Raiz e com que número.")
    if outros:
        total = len(h["inventoryIds"])
        if total > 1 and len(outros) < total:
            partes.append(f"Dos {total} números, só {q(outros)} não existe no Raiz; os outros estão certos. "
                          f"Confirmar o número dessa peça e corrigi-lo {onde}.")
        else:
            partes.append(f"O número {q(outros)} não existe no Raiz. "
                          f"Confirmar o número de inventário desta peça no Raiz e corrigi-lo {onde}.")
    same = by_title.get(norm_title(re.sub(r"^\d+\.\s*", "", h["title"] or "")), [])
    if 0 < len(same) <= 2 and (outros or dep):
        partes.append(f"No Raiz existe «{clean_html(h['title'])}» com o número {q(same)} — confirmar se é a mesma peça.")
    if not h["image"] and h.get("media") and all(m.lower().endswith(".glb") for m in h["media"] if m):
        partes.append("(Sem fotografia na visita, só modelo 3D.)")
    return " ".join(partes)


# ---------------------------------------------------------------------- excel
def write_table(out_path: Path, title: str, headers, widths, rows, images=None, thumbs_dir: Path | None = None):
    """Um Excel simples: cabecalho escuro, texto com wrap, opcionalmente uma coluna A com miniaturas.
    `images` e uma lista paralela a `rows` com o caminho da imagem (ou None) de cada linha."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

    wb = Workbook()
    ws = wb.active
    ws.title = title
    ws.append(headers)
    thin = Side(style="thin", color="BBBBBB")
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="444444")
        c.alignment = Alignment(wrap_text=True, vertical="center")
    for i, w in enumerate(widths):
        ws.column_dimensions[chr(65 + i)].width = w
    if images:
        from PIL import Image as PImage
        from openpyxl.drawing.image import Image as XLImage
        thumbs_dir.mkdir(parents=True, exist_ok=True)
    for r, row in enumerate(rows, start=2):
        ws.append(list(row))
        for c in ws[r]:
            c.alignment = Alignment(wrap_text=True, vertical="top")
            c.border = Border(top=thin, bottom=thin, left=thin, right=thin)
        if images:
            ws.row_dimensions[r].height = 150
            if images[r - 2]:
                im = PImage.open(images[r - 2]).convert("RGB")
                im.thumbnail((230, 190))
                t = thumbs_dir / f"{out_path.stem}_{r}.png"
                im.save(t)
                ws.add_image(XLImage(str(t)), f"A{r}")
    ws.freeze_panes = "B2" if images else "A2"
    for cand in [out_path] + [out_path.with_name(out_path.stem + f"_v{i}" + out_path.suffix) for i in range(2, 6)]:
        try:
            wb.save(cand)
            return cand
        except PermissionError:
            continue
    raise SystemExit(f"Nao consegui gravar {out_path}: fecha o ficheiro no Excel.")


# ----------------------------------------------------------------------- main
def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("path", help=".vtp do projecto ou pasta exportada da visita")
    ap.add_argument("--museum", required=True, help="museum_id no indice (mnt, mnaz, mnsr, mj, mnc, ...)")
    ap.add_argument("--index", help="indice de artefactos (default: o do .env do backend)")
    ap.add_argument("--fix", action="store_true", help="corrigir no ficheiro o que for mecanico (faz .bak)")
    ap.add_argument("--location-from", choices=["subtitle", "label"],
                    help="com --fix: o location de cada hotspot passa a ser o subtitle (ou label) do panorama onde esta")
    ap.add_argument("--images-dir", help="pasta extra com imagens (ex.: renders dos .glb), nome = ficheiro do modelo")
    ap.add_argument("--out-dir", default=str(REPO / "FIXES"), help="pasta base dos resultados; cria <pasta>/<museu>/ (default: FIXES/)")
    ap.add_argument("--no-xlsx", action="store_true", help="so validar, sem gerar o Excel")
    a = ap.parse_args()
    path = Path(a.path).resolve()
    out_dir = Path(a.out_dir).resolve() / a.museum   # uma pasta por museu
    images_dir = Path(a.images_dir).resolve() if a.images_dir else None
    is_vtp = path.is_file() and path.suffix.lower() == ".vtp"

    client, default_index = connect()
    index = a.index or default_index
    catalogue, museum_name = load_catalogue(client, index, a.museum)
    if not catalogue:
        sys.exit(f"O indice {index} nao tem nenhuma peca com museum_id={a.museum!r}.")
    by_norm, by_title = {}, {}
    for inv, title in catalogue.items():
        by_norm.setdefault(norm(inv), inv)
        by_title.setdefault(norm_title(title), []).append(inv)

    load = hotspots_from_vtp if is_vtp else hotspots_from_export
    hotspots = load(path)
    print(f"indice   : {index}  (museum_id={a.museum}, {len(catalogue)} pecas, {museum_name})")

    # o plano e calculado sempre: com --fix aplica-se, sem --fix so vai para o Excel dos «mal escritos»
    plan = plan_fixes(hotspots, catalogue, by_norm, a.museum, location_from=a.location_from)
    before = {h["key"]: h for h in hotspots if h["hasHotspot"]}
    if a.fix:
        if is_vtp:
            # mesmo sem plano para os hotspots, os URLs das janelas podem precisar de correccao
            n, urls, b = apply_fix_vtp(path, plan)
            if b:
                print(f"corrigido: {n} hotspots, {urls} URLs de janela  (copia: {b.name})")
            else:
                print("corrigido: nada a corrigir mecanicamente")
        elif plan:
            n, bs = apply_fix_export(path, plan)
            print(f"corrigido: {n} blocos em {len(bs)} ficheiro(s)  (copias: {', '.join(x.name for x in bs)})")
        else:
            print("corrigido: nada a corrigir mecanicamente")
        for key, p in plan.items():
            print(f"  - {p['title'][:44]:44} {'; '.join(p['changes'])}")
        if plan or (is_vtp and b):
            hotspots = load(path)

    if images_dir:
        for h in hotspots:
            if not h["image"]:
                for m in h.get("media", []):
                    c = images_dir / (os.path.splitext(m)[0] + ".jpg")
                    if c.exists():
                        h["image"] = str(c); break

    dominant = find_dominant_prefix([i for h in hotspots for i in h["inventoryIds"]], by_norm)
    ok, formato, falta, sn, rows, sn_nome = [], [], [], [], [], []
    for h in hotspots:
        if not h["hasHotspot"]:
            # ficha ainda sem hotspot: o que interessa e o numero que se vai escrever, por isso tolera prefixo/caixa
            h["rawInventory"] = ", ".join(h["inventoryIds"])
            resolved = []
            for raw in h["inventoryIds"]:
                real = lookup(raw, by_norm, dominant)
                parts = None if real else split_glued(raw, by_norm, dominant)
                resolved += [real] if real else (parts or [strip_prefix(raw, dominant)])
            h["inventoryIds"] = resolved
        # sem --fix, um numero so mal escrito («MMC_Inv.A4121») nao e «nao existe»: o que se
        # avalia contra o Raiz e o numero como ficaria depois da correccao (o plano)
        planned = (plan.get(h["key"]) or {}).get("ids") if h["hasHotspot"] else None
        eval_ids = planned or h["inventoryIds"]
        # um id SN (MUSEU_SN_n ou MUSEU_SN_Nome) e tratado como qualquer numero: so e valido
        # se ja estiver no indice; senao vai para «em falta», para ser indexado a mao
        missing = []
        for inv in h["inventoryIds"]:
            is_sn = bool(SN_OK_RE.match(inv) or SN_NAMED_RE.match(inv))
            if inv in catalogue:
                (sn if is_sn else ok).append((h, inv)); continue
            if is_sn:
                sn_nome.append((h, inv)); continue
            real = by_norm.get(norm(inv))
            if real:
                formato.append((h, inv, real))
        for inv in eval_ids:
            if inv in catalogue or by_norm.get(norm(inv)):
                continue
            falta.append((h, inv)); missing.append(inv)
        if missing:
            rows.append({"image": h["image"], "nome": clean_html(h["title"]), "ficha": h["ficha"], "loc": h["location"],
                         "pano": h["panorama"] if h["hasHotspot"] else "— (sem hotspot)",
                         "inv": ", ".join(eval_ids),
                         "nota": nota_para(dict(h, inventoryIds=eval_ids), missing, museum_name or a.museum.upper(), by_title, by_norm)})

    com_hotspot = [h for h in hotspots if h["hasHotspot"]]
    print(f"hotspots : {len(com_hotspot)}   fichas sem hotspot: {len(hotspots) - len(com_hotspot)}")
    print(f"  numeros certos            : {len(ok)}")
    print(f"  existem, formato diferente: {len(formato)}")
    print(f"  NAO existem no Raiz       : {len(falta)}")
    print(f"  SN ja no indice           : {len(sn)}")
    curvas = [h for h in hotspots if h["smartQuotes"] and not re.search(r"[“”]", h["title"] or "")]
    if curvas:
        print(f"  blocos com aspas curvas   : {len(curvas)}")
    stale = [h for h in hotspots if h.get("sentencesStale")]
    if stale:
        print(f"  sentences desatualizado   : {len(stale)}  (o export do 3DVista sai com o codigo antigo)")
    # um id SN (com numero ou com nome) e valido desde que seja unico: o mesmo id em dois hotspots
    # tem de ser a mesma peca (vista de dois panoramas), nunca duas pecas diferentes
    sn_titles = {}
    for h in hotspots:
        for i in ((plan.get(h["key"]) or {}).get("ids") or h["inventoryIds"]) if h["hasHotspot"] else []:
            if SN_OK_RE.match(i) or SN_NAMED_RE.match(i):
                sn_titles.setdefault(i, set()).add(norm_title(h["title"]))
    sn_conflict = {i for i, t in sn_titles.items() if len(t) > 1}
    if sn_nome:
        print(f"  SN por indexar            : {len(sn_nome)}  (em falta ate serem indexados a mao)")
    if sn_conflict:
        print(f"  SN repetido em pecas diferentes: {len(sn_conflict)}  {sorted(sn_conflict)}")

    def line(h, extra):
        return f'  {str(h["panorama"] or "(sem hotspot)")[:22]:22} {str(h["title"] or "")[:46]:46} {extra}'

    if formato:
        print("\n== EXISTEM MAS ESCRITOS DE FORMA DIFERENTE (escrever como no indice) ==")
        for h, inv, real in formato:
            print(line(h, f"{inv!r}  ->  {real!r}"))
    falta_raiz = [(h, inv) for h, inv in falta if not (SN_OK_RE.match(inv) or SN_NAMED_RE.match(inv))]
    if falta_raiz:
        print("\n== NAO EXISTEM NO RAIZ ==")
        for h, inv in falta_raiz:
            print(line(h, repr(inv)))
    if curvas:
        print("\n== ASPAS CURVAS (o postMessage nunca dispara) ==")
        for h in curvas:
            print(line(h, ""))
    if sn_nome:
        print("\n== SN POR INDEXAR (nao estao no indice; indexar a mao com este id) ==")
        for h, inv in sn_nome:
            print(line(h, repr(inv)))
    if sn_conflict:
        print("\n== SN REPETIDO EM PECAS DIFERENTES (tem de ser um id por peca) ==")
        for h in hotspots:
            for i in h["inventoryIds"]:
                if i in sn_conflict:
                    print(line(h, repr(i)))

    out_dir.mkdir(parents=True, exist_ok=True)
    (out_dir / "validacao.json").write_text(json.dumps({
        "index": index, "museum": a.museum, "hotspots": hotspots,
        "ok": [inv for _, inv in ok], "formato": [(inv, real) for _, inv, real in formato],
        "falta": [inv for _, inv in falta], "sn": [inv for _, inv in sn],
        "mal_escritos": {k: {"changes": p["changes"], "ids": p["ids"], "location": p["location"]} for k, p in plan.items()},
    }, ensure_ascii=False, indent=1), encoding="utf-8")
    print(f"\njson  : {out_dir / 'validacao.json'}")

    if not a.no_xlsx:
        thumbs = out_dir / ".thumbs"
        # 1) numeros que nao existem no Raiz (hotspots e fichas) — e o que se envia ao museu
        p1 = write_table(out_dir / "nao_existem_no_raiz.xlsx", "Não existem no Raiz",
                         ["Imagem", "Nome", "Ficha (na visita)", "Location", "Panorama no 3DVista", "Nº de inventário na visita", "Nota"],
                         [34, 30, 46, 24, 22, 26, 60],
                         [["", r["nome"], r["ficha"], r["loc"], r["pano"], r["inv"], r["nota"]] for r in rows],
                         images=[r["image"] for r in rows], thumbs_dir=thumbs)
        print(f"excel : {p1}  ({len(rows)} pecas, {sum(1 for r in rows if not r['image'])} sem imagem)")

        # 2) fichas em Conteudos sem hotspot nenhum — e o que se envia a quem faz a visita
        orphans = []
        for h in hotspots:
            if h["hasHotspot"]:
                continue
            exists = [i in catalogue for i in h["inventoryIds"]]
            orphans.append([h["location"], clean_html(h["title"]), h.get("rawInventory", ""), ", ".join(h["inventoryIds"]),
                            "Sim" if exists and all(exists) else ("Em parte" if any(exists) else "Não")])
        p2 = write_table(out_dir / "sem_hotspot.xlsx", "Sem hotspot",
                         ["Sala", "Ficha", "Nº na ficha", "Nº a escrever no hotspot", "Existe no Raiz?"],
                         [24, 44, 34, 44, 16], orphans)
        print(f"excel : {p2}  ({len(orphans)} fichas)")

        # 4) hotspots validos: todos os numeros existem no Raiz (ou sao vitrines MUSEU_SN_n)
        validos = []
        for h in hotspots:
            if not h["hasHotspot"]:
                continue
            ids = (plan.get(h["key"]) or {}).get("ids") or h["inventoryIds"]
            is_sn = lambda i: bool(SN_OK_RE.match(i) or SN_NAMED_RE.match(i))
            if not ids or not all(i in catalogue for i in ids) or any(i in sn_conflict for i in ids):
                continue
            raiz = "; ".join(catalogue.get(i) or "" for i in ids)
            validos.append([h["panorama"], clean_html(h["title"]), ", ".join(ids), raiz,
                            (plan.get(h["key"]) or {}).get("location") or h["location"],
                            "Sim" if any(is_sn(i) for i in ids) else "Não"])
        p4 = write_table(out_dir / "validos.xlsx", "Válidos",
                         ["Panorama", "Peça (no hotspot)", "Nº de inventário", "Título no índice", "Location", "SN"],
                         [22, 44, 30, 44, 26, 8], validos)
        print(f"excel : {p4}  ({len(validos)} hotspots válidos)")

        # 3) hotspots mal escritos: formato, colados, SN, aspas, HTML, location — o que o --fix corrige
        mal = []
        for key, p in plan.items():
            h = before.get(key, {})
            mal.append([h.get("panorama"), clean_html(h.get("title") or ""), ", ".join(h.get("inventoryIds") or []),
                        h.get("location"), "; ".join(p["changes"]), ", ".join(p["ids"] or []), p["location"],
                        "Sim" if a.fix else "Não (correr com --fix)"])
        p3 = write_table(out_dir / "mal_escritos.xlsx", "Mal escritos",
                         ["Panorama", "Peça", "Nº no hotspot (antes)", "Location (antes)", "Problemas", "Nº como deve ficar", "Location como deve ficar", "Corrigido automaticamente"],
                         [22, 40, 30, 24, 60, 34, 26, 18], mal)
        print(f"excel : {p3}  ({len(mal)} hotspots)")
    sys.exit(1 if falta else 0)


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    main()
